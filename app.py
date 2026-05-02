import streamlit as st
import pandas as pd
import numpy as np
from io import BytesIO
import os
from pathlib import Path
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import xlrd
import warnings

warnings.filterwarnings("ignore")

st.set_page_config(
    page_title="Excel Consolidator",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── Custom CSS ────────────────────────────────────────────────────────────────
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');

    html, body, [class*="css"] { font-family: 'Inter', sans-serif; }

    .main { background: #f8fafc; }

    .hero {
        background: linear-gradient(135deg, #1e3a5f 0%, #2d6a9f 50%, #1a8cff 100%);
        border-radius: 16px;
        padding: 2.5rem 2rem;
        margin-bottom: 1.5rem;
        color: white;
        box-shadow: 0 8px 32px rgba(30,58,95,0.18);
    }
    .hero h1 { font-size: 2.1rem; font-weight: 700; margin: 0 0 0.3rem; letter-spacing: -0.5px; }
    .hero p  { font-size: 1rem; opacity: 0.88; margin: 0; }

    .stat-card {
        background: white;
        border-radius: 12px;
        padding: 1.2rem 1.5rem;
        box-shadow: 0 2px 12px rgba(0,0,0,0.07);
        border-left: 4px solid #2d6a9f;
        margin-bottom: 0.5rem;
    }
    .stat-card .label { font-size: 0.78rem; color: #64748b; font-weight: 500; text-transform: uppercase; letter-spacing: 0.05em; }
    .stat-card .value { font-size: 1.6rem; font-weight: 700; color: #1e3a5f; margin-top: 0.1rem; }
    .stat-card .sub   { font-size: 0.82rem; color: #94a3b8; }

    .file-badge {
        display: inline-block;
        padding: 2px 10px;
        border-radius: 20px;
        font-size: 0.75rem;
        font-weight: 600;
        margin: 2px;
    }
    .badge-html { background: #dbeafe; color: #1d4ed8; }
    .badge-xls  { background: #dcfce7; color: #166534; }
    .badge-xlsx { background: #fef9c3; color: #854d0e; }
    .badge-csv  { background: #fce7f3; color: #9d174d; }

    .step-box {
        background: white;
        border-radius: 10px;
        padding: 1rem 1.2rem;
        margin-bottom: 0.8rem;
        box-shadow: 0 1px 6px rgba(0,0,0,0.06);
        border: 1px solid #e2e8f0;
    }
    .step-num {
        display: inline-block;
        width: 26px; height: 26px;
        background: #2d6a9f;
        color: white;
        border-radius: 50%;
        text-align: center;
        line-height: 26px;
        font-size: 0.8rem;
        font-weight: 700;
        margin-right: 8px;
    }

    .success-box {
        background: #f0fdf4;
        border: 1.5px solid #86efac;
        border-radius: 10px;
        padding: 1rem 1.2rem;
        margin: 1rem 0;
    }
    .warning-box {
        background: #fffbeb;
        border: 1.5px solid #fcd34d;
        border-radius: 10px;
        padding: 0.8rem 1.2rem;
        margin: 0.5rem 0;
    }
    .error-box {
        background: #fef2f2;
        border: 1.5px solid #fca5a5;
        border-radius: 10px;
        padding: 0.8rem 1.2rem;
        margin: 0.5rem 0;
    }

    div[data-testid="stFileUploader"] {
        border: 2px dashed #93c5fd;
        border-radius: 12px;
        padding: 0.5rem;
        background: #eff6ff;
    }

    .stButton > button {
        background: linear-gradient(135deg, #1e3a5f, #2d6a9f);
        color: white;
        border: none;
        border-radius: 8px;
        padding: 0.55rem 1.5rem;
        font-weight: 600;
        font-size: 0.95rem;
        transition: all 0.2s;
        box-shadow: 0 2px 8px rgba(30,58,95,0.2);
    }
    .stButton > button:hover { transform: translateY(-1px); box-shadow: 0 4px 14px rgba(30,58,95,0.3); }

    .stDownloadButton > button {
        background: linear-gradient(135deg, #166534, #16a34a) !important;
        color: white !important;
        border: none !important;
        border-radius: 8px !important;
        font-weight: 600 !important;
        width: 100%;
    }

    [data-testid="stSidebar"] { background: #1e3a5f; }
    [data-testid="stSidebar"] * { color: #e2e8f0 !important; }
    [data-testid="stSidebar"] .stSelectbox label,
    [data-testid="stSidebar"] .stCheckbox label { color: #cbd5e1 !important; }

    .dataframe thead tr th {
        background: #1e3a5f !important;
        color: white !important;
        font-weight: 600 !important;
    }
</style>
""", unsafe_allow_html=True)


# ── Helpers ───────────────────────────────────────────────────────────────────

def detect_format(file_bytes: bytes, filename: str) -> str:
    """Detect if file is real XLS (OLE2), HTML-as-XLS, XLSX, or CSV."""
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        return "csv"
    if ext in (".xlsx", ".xlsm"):
        return "xlsx"
    # Peek at magic bytes
    if file_bytes[:4] == b"\xd0\xcf\x11\xe0":
        return "xls_real"
    # Likely HTML disguised as XLS
    snippet = file_bytes[:200].decode("utf-8", errors="ignore").lower()
    if "<table" in snippet or "<html" in snippet or "<style" in snippet:
        return "xls_html"
    # Try anyway
    return "xls_real"


def read_html_xls(file_bytes: bytes, source_name: str) -> tuple[pd.DataFrame, str]:
    content = file_bytes.decode("utf-8", errors="ignore")
    soup = BeautifulSoup(content, "html.parser")
    table = soup.find("table")
    if not table:
        raise ValueError("No HTML table found in file.")
    header_row = table.find("tr")
    headers = [th.get_text(strip=True) for th in header_row.find_all(["th", "td"])]
    rows = []
    for tr in table.find_all("tr")[1:]:
        cells = [td.get_text(separator=" ", strip=True) for td in tr.find_all("td")]
        if any(c.strip() for c in cells):
            rows.append(cells)
    if not rows:
        raise ValueError("Table found but no data rows.")
    ncols = max(len(r) for r in rows)
    headers = (headers + [""] * ncols)[:ncols]
    df = pd.DataFrame(rows, columns=headers)
    return df, "HTML-as-XLS"


def read_real_xls(file_bytes: bytes, source_name: str, sheet_choice: str | None) -> tuple[pd.DataFrame, str]:
    bio = BytesIO(file_bytes)
    sheets = pd.read_excel(bio, engine="xlrd", sheet_name=None, dtype=str)
    if not sheets:
        raise ValueError("No sheets found.")
    if sheet_choice and sheet_choice in sheets:
        df = sheets[sheet_choice]
    else:
        df = list(sheets.values())[0]
    df = df.dropna(how="all").reset_index(drop=True)
    return df, "XLS (Binary)"


def read_xlsx(file_bytes: bytes, source_name: str, sheet_choice: str | None) -> tuple[pd.DataFrame, str]:
    bio = BytesIO(file_bytes)
    sheets = pd.read_excel(bio, engine="openpyxl", sheet_name=None, dtype=str)
    if not sheets:
        raise ValueError("No sheets found.")
    if sheet_choice and sheet_choice in sheets:
        df = sheets[sheet_choice]
    else:
        df = list(sheets.values())[0]
    df = df.dropna(how="all").reset_index(drop=True)
    return df, "XLSX"


def read_csv(file_bytes: bytes, source_name: str, encoding: str = "utf-8") -> tuple[pd.DataFrame, str]:
    for enc in [encoding, "latin-1", "cp1252"]:
        try:
            df = pd.read_csv(BytesIO(file_bytes), dtype=str, encoding=enc)
            df = df.dropna(how="all").reset_index(drop=True)
            return df, "CSV"
        except Exception:
            continue
    raise ValueError("Could not decode CSV file.")


def get_sheet_names(file_bytes: bytes, filename: str) -> list[str]:
    ext = Path(filename).suffix.lower()
    try:
        if ext in (".xlsx", ".xlsm"):
            bio = BytesIO(file_bytes)
            wb = pd.ExcelFile(bio, engine="openpyxl")
            return wb.sheet_names
        elif ext in (".xls",):
            fmt = detect_format(file_bytes, filename)
            if fmt == "xls_real":
                bio = BytesIO(file_bytes)
                wb = pd.ExcelFile(bio, engine="xlrd")
                return wb.sheet_names
    except Exception:
        pass
    return []


def parse_file(file_bytes: bytes, filename: str, sheet_choice: str | None = None) -> tuple[pd.DataFrame, str]:
    fmt = detect_format(file_bytes, filename)
    if fmt == "csv":
        return read_csv(file_bytes, filename)
    elif fmt == "xlsx":
        return read_xlsx(file_bytes, filename, sheet_choice)
    elif fmt == "xls_html":
        return read_html_xls(file_bytes, filename)
    else:
        return read_real_xls(file_bytes, filename, sheet_choice)


def build_excel_output(combined: pd.DataFrame, include_summary: bool) -> bytes:
    wb = Workbook()

    H_FILL   = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
    H_FONT   = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    D_FONT   = Font(name="Arial", size=9)
    ALT_FILL = PatternFill("solid", start_color="EBF3FB", end_color="EBF3FB")
    SRC_FILL = PatternFill("solid", start_color="FFF9C4", end_color="FFF9C4")
    SRC_FONT = Font(name="Arial", size=9, color="7B4F00")
    CENTER   = Alignment(horizontal="center", vertical="center")
    LEFT     = Alignment(horizontal="left",   vertical="center")
    thin     = Side(style="thin", color="CCCCCC")
    BORDER   = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── Sheet 1: All Data ──────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Consolidated Data"
    cols = list(combined.columns)

    for ci, col in enumerate(cols, 1):
        cell = ws.cell(row=1, column=ci, value=col)
        cell.font   = H_FONT
        cell.fill   = H_FILL
        cell.alignment = CENTER
        cell.border = BORDER
        if col == "Source File":
            cell.fill = PatternFill("solid", start_color="B8860B", end_color="B8860B")

    for ri, row_data in enumerate(combined.itertuples(index=False), 2):
        alt = ri % 2 == 0
        for ci, val in enumerate(row_data, 1):
            v = "" if (val is None or (isinstance(val, float) and np.isnan(val))) else str(val)
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.border = BORDER
            cell.font   = D_FONT
            col_name = cols[ci - 1]
            if col_name == "Source File":
                cell.fill      = SRC_FILL
                cell.font      = SRC_FONT
                cell.alignment = CENTER
            else:
                cell.fill      = ALT_FILL if alt else PatternFill()
                cell.alignment = CENTER if ci <= 3 else LEFT

    # Auto-width (capped)
    for ci, col in enumerate(cols, 1):
        max_len = len(str(col))
        for ri in range(2, min(len(combined) + 2, 200)):
            v = ws.cell(row=ri, column=ci).value or ""
            max_len = max(max_len, len(str(v)))
        ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 3, 45)

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "B2"

    # ── Sheet 2: File Summary ──────────────────────────────────────────────
    if include_summary and "Source File" in combined.columns:
        ws2 = wb.create_sheet("File Summary")
        sum_hdrs = ["Source File", "Rows", "Columns"]
        for ci, h in enumerate(sum_hdrs, 1):
            cell = ws2.cell(row=1, column=ci, value=h)
            cell.font = H_FONT; cell.fill = H_FILL
            cell.alignment = CENTER; cell.border = BORDER

        groups = combined.groupby("Source File", sort=False)
        for ri, (src, grp) in enumerate(groups, 2):
            alt = ri % 2 == 0
            for ci, val in enumerate([src, len(grp), len(combined.columns) - 1], 1):
                cell = ws2.cell(row=ri, column=ci, value=val)
                cell.font = D_FONT; cell.alignment = CENTER; cell.border = BORDER
                if alt: cell.fill = ALT_FILL

        tr = len(groups) + 2
        for ci, val in enumerate(["TOTAL", len(combined), ""], 1):
            cell = ws2.cell(row=tr, column=ci, value=val)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
            cell.fill = H_FILL; cell.alignment = CENTER; cell.border = BORDER

        for ci, w in enumerate([40, 12, 12], 1):
            ws2.column_dimensions[get_column_letter(ci)].width = w
        ws2.row_dimensions[1].height = 22

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Session state ─────────────────────────────────────────────────────────────
if "parsed_files" not in st.session_state:
    st.session_state.parsed_files   = {}   # filename -> df
    st.session_state.file_formats   = {}   # filename -> format string
    st.session_state.file_errors    = {}   # filename -> error string
    st.session_state.sheet_choices  = {}   # filename -> chosen sheet
    st.session_state.combined_df    = None


# ── Sidebar ───────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("## ⚙️ Options")
    st.markdown("---")

    output_filename = st.text_input("Output filename", value="Consolidated_Output")
    include_summary = st.checkbox("Add File Summary sheet", value=True)

    st.markdown("### 🧹 Column Handling")
    col_strategy = st.selectbox(
        "Mismatched columns",
        ["Union (keep all columns)", "Intersection (common columns only)"],
    )

    st.markdown("### 🔢 Row Numbering")
    add_row_num = st.checkbox("Add global row number column", value=False)

    st.markdown("---")
    st.markdown("### 📋 Supported Formats")
    for fmt, badge in [
        ("`.xls` — Binary XLS", "badge-xls"),
        ("`.xls` — HTML-as-XLS", "badge-html"),
        ("`.xlsx` / `.xlsm`", "badge-xlsx"),
        ("`.csv`", "badge-csv"),
    ]:
        st.markdown(f'<span class="file-badge {badge}">{fmt}</span>', unsafe_allow_html=True)


# ── Hero ──────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="hero">
    <h1>📊 Excel File Consolidator</h1>
    <p>Merge multiple Excel / CSV files of any format into one clean, formatted spreadsheet — with a <strong>Source File</strong> column tracking every row's origin.</p>
</div>
""", unsafe_allow_html=True)


# ── Step 1: Upload ─────────────────────────────────────────────────────────────
st.markdown('<div class="step-box"><span class="step-num">1</span> <b>Upload your files</b></div>', unsafe_allow_html=True)

uploaded = st.file_uploader(
    "Drop files here (XLS, XLSX, CSV — any mix)",
    type=["xls", "xlsx", "xlsm", "csv"],
    accept_multiple_files=True,
    label_visibility="collapsed",
)

if uploaded:
    st.markdown(f"**{len(uploaded)} file(s) selected**")

    # Sheet selector for multi-sheet files
    sheet_selections = {}
    for uf in uploaded:
        file_bytes = uf.read()
        uf.seek(0)
        sheets = get_sheet_names(file_bytes, uf.name)
        if len(sheets) > 1:
            with st.expander(f"📋 `{uf.name}` has {len(sheets)} sheets — choose one"):
                chosen = st.selectbox(
                    f"Sheet for {uf.name}",
                    sheets,
                    key=f"sheet_{uf.name}",
                    label_visibility="collapsed",
                )
                sheet_selections[uf.name] = chosen

    # ── Step 2: Parse ──────────────────────────────────────────────────────────
    st.markdown('<div class="step-box"><span class="step-num">2</span> <b>Parse & preview files</b></div>', unsafe_allow_html=True)

    if st.button("🔍 Parse All Files"):
        st.session_state.parsed_files  = {}
        st.session_state.file_formats  = {}
        st.session_state.file_errors   = {}
        st.session_state.combined_df   = None

        progress = st.progress(0)
        status   = st.empty()

        for i, uf in enumerate(uploaded):
            status.text(f"Parsing {uf.name}…")
            try:
                file_bytes = uf.read()
                sheet_ch   = sheet_selections.get(uf.name)
                df, fmt    = parse_file(file_bytes, uf.name, sheet_ch)
                df = df.dropna(how="all").reset_index(drop=True)
                st.session_state.parsed_files[uf.name]  = df
                st.session_state.file_formats[uf.name]  = fmt
            except Exception as e:
                st.session_state.file_errors[uf.name] = str(e)
            progress.progress((i + 1) / len(uploaded))

        status.empty()
        progress.empty()

    # Show parse results
    if st.session_state.parsed_files:
        for fname, df in st.session_state.parsed_files.items():
            fmt = st.session_state.file_formats.get(fname, "")
            badge_cls = {"HTML-as-XLS": "badge-html", "XLS (Binary)": "badge-xls",
                         "XLSX": "badge-xlsx", "CSV": "badge-csv"}.get(fmt, "badge-xls")
            with st.expander(f"✅  `{fname}` — {len(df):,} rows × {len(df.columns)} cols"):
                st.markdown(f'<span class="file-badge {badge_cls}">{fmt}</span>', unsafe_allow_html=True)
                st.dataframe(df.head(5), use_container_width=True, hide_index=True)

    for fname, err in st.session_state.file_errors.items():
        st.markdown(f'<div class="error-box">❌ <b>{fname}</b>: {err}</div>', unsafe_allow_html=True)

    # ── Step 3: Consolidate ────────────────────────────────────────────────────
    if st.session_state.parsed_files:
        st.markdown('<div class="step-box"><span class="step-num">3</span> <b>Consolidate</b></div>', unsafe_allow_html=True)

        if st.button("⚡ Consolidate All Files"):
            dfs = []
            for fname, df in st.session_state.parsed_files.items():
                df = df.copy()
                df.insert(0, "Source File", fname)
                dfs.append(df)

            # Column alignment
            if col_strategy.startswith("Union"):
                combined = pd.concat(dfs, ignore_index=True, sort=False)
            else:
                common = set(dfs[0].columns)
                for d in dfs[1:]:
                    common &= set(d.columns)
                common = sorted(common, key=lambda c: list(dfs[0].columns).index(c) if c in dfs[0].columns else 999)
                combined = pd.concat([d[list(common)] for d in dfs], ignore_index=True, sort=False)

            if add_row_num:
                combined.insert(1, "Row No.", range(1, len(combined) + 1))

            st.session_state.combined_df = combined

        if st.session_state.combined_df is not None:
            combined = st.session_state.combined_df

            # Stats
            file_count = combined["Source File"].nunique() if "Source File" in combined.columns else len(st.session_state.parsed_files)
            c1, c2, c3, c4 = st.columns(4)
            for col_obj, label, value, sub in [
                (c1, "Total Rows",    f"{len(combined):,}",         "across all files"),
                (c2, "Total Columns", f"{len(combined.columns):,}", "in output sheet"),
                (c3, "Files Merged",  f"{file_count}",              "source files"),
                (c4, "Output Size",   f"~{len(combined) * len(combined.columns) // 1000}K", "cells"),
            ]:
                col_obj.markdown(f"""
                <div class="stat-card">
                    <div class="label">{label}</div>
                    <div class="value">{value}</div>
                    <div class="sub">{sub}</div>
                </div>""", unsafe_allow_html=True)

            st.markdown("**Preview — first 50 rows**")
            st.dataframe(combined.head(50), use_container_width=True, hide_index=True)

            # ── Step 4: Download ───────────────────────────────────────────────
            st.markdown('<div class="step-box"><span class="step-num">4</span> <b>Download</b></div>', unsafe_allow_html=True)

            dl_col1, dl_col2 = st.columns(2)

            with dl_col1:
                with st.spinner("Building Excel file…"):
                    xlsx_bytes = build_excel_output(combined, include_summary)
                st.download_button(
                    label="⬇️ Download as Excel (.xlsx)",
                    data=xlsx_bytes,
                    file_name=f"{output_filename}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )

            with dl_col2:
                csv_bytes = combined.to_csv(index=False).encode("utf-8-sig")
                st.download_button(
                    label="⬇️ Download as CSV",
                    data=csv_bytes,
                    file_name=f"{output_filename}.csv",
                    mime="text/csv",
                )

else:
    # Empty state
    st.markdown("""
    <div style="text-align:center; padding: 3rem 1rem; color: #94a3b8;">
        <div style="font-size: 3.5rem; margin-bottom: 1rem;">📂</div>
        <div style="font-size: 1.1rem; font-weight: 500;">Upload your Excel or CSV files above to get started</div>
        <div style="font-size: 0.9rem; margin-top: 0.5rem;">Supports XLS (binary & HTML), XLSX, XLSM, and CSV</div>
    </div>
    """, unsafe_allow_html=True)

# ── Footer ─────────────────────────────────────────────────────────────────────
st.markdown("---")
st.markdown(
    '<p style="text-align:center; color:#94a3b8; font-size:0.8rem;">Excel Consolidator · Handles XLS (binary & HTML-as-XLS), XLSX, XLSM, CSV · Source column auto-added</p>',
    unsafe_allow_html=True,
)
